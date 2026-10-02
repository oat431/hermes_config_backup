#!/usr/bin/env python3
"""Convert a markdown file (Thai or mixed) into a forward-friendly .docx.

Usage: uv run --with python-docx python md_to_docx.py <src.md> <out.docx> [font]
Default font: Leelawadee UI (Windows Thai UI font; Sarabun/Tahoma are OK fallbacks).

Handles: h1-h3 headings, blockquotes ('> '), bullet/numbered lists, **bold**, `code`,
horizontal rules (skipped), and skips lone '>' separator lines (which would otherwise
render as stray '>' paragraphs). Sets the complex-script font attribute so Thai text
renders with the intended font.
"""
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def main():
    src, out = sys.argv[1], sys.argv[2]
    font = sys.argv[3] if len(sys.argv) > 3 else "Leelawadee UI"

    doc = Document()

    def set_font(style, name):
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)

    set_font(doc.styles["Normal"], font)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            set_font(doc.styles[name], font)
        except KeyError:
            pass

    def add_runs(p, text):
        for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                p.add_run(part[1:-1]).font.name = "Consolas"
            else:
                p.add_run(part)

    for line in open(src, encoding="utf-8").read().splitlines():
        l = line.rstrip()
        if not l.strip() or l.startswith("---") or l.strip() == ">":
            continue
        if l.startswith("# "):
            doc.add_heading(l[2:], level=0)
            continue
        if l.startswith("## "):
            doc.add_heading(l[3:], level=1)
            continue
        if l.startswith("### "):
            doc.add_heading(l[4:], level=2)
            continue
        if l.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, l[2:])
            for r in p.runs:
                r.italic = True
            continue
        m = re.match(r"^\s*[-*] (.*)$", l)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\s*\d+\. (.*)$", l)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        add_runs(doc.add_paragraph(), l)

    doc.save(out)
    print("WROTE", out)


if __name__ == "__main__":
    main()
